# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

test_cases = [
    # ===== 模块1: 任务创建 - 复刻视频面板 =====
    {
        "id": "TC-001",
        "module": "任务创建-面板切换",
        "name": "新建任务弹窗-切换到复刻视频面板",
        "priority": "P0",
        "precondition": "1. 用户已登录系统\n2. 点击「新建任务」打开弹窗",
        "steps": "1. 确认弹窗默认显示「视频解析任务」面板\n2. 点击「复刻视频」TAB\n3. 观察面板内容是否切换",
        "expected": "1. 弹窗标题切换为「新增复刻视频任务」\n2. 面板展示复刻视频表单：视频名称、复刻内容描述、视频上传方式（链接/文件上传）\n3. 顶部提示文案显示：视频格式仅支持.mp4、文件最大100MB、视频最长30秒，每个片段4-10秒、消耗5积分",
        "requirement": "复刻视频-面板切换"
    },
    {
        "id": "TC-002",
        "module": "任务创建-面板切换",
        "name": "新建任务弹窗-TAB间来回切换",
        "priority": "P1",
        "precondition": "新建任务弹窗已打开",
        "steps": "1. 点击「复刻视频」TAB\n2. 切换到「AI自动生成脚本」TAB\n3. 再切回「复刻视频」TAB\n4. 观察表单状态",
        "expected": "1. 每次切换面板内容正确对应\n2. 复刻视频面板的输入内容切换后不丢失（如已输入视频名称）\n3. TAB高亮状态正确切换",
        "requirement": "复刻视频-面板切换"
    },

    # ===== 模块2: 视频上传 - 文件上传校验 =====
    {
        "id": "TC-003",
        "module": "文件上传校验",
        "name": "文件上传方式-选择有效.mp4文件（正常流程）",
        "priority": "P0",
        "precondition": "1. 复刻视频面板已打开\n2. 选择「文件上传」方式",
        "steps": "1. 点击「文件上传」单选按钮\n2. 观察文件上传区域显示\n3. 点击「选择文件」按钮\n4. 选择一个符合要求的.mp4文件（≤100MB，≤30秒）\n5. 观察文件名显示和后续流程",
        "expected": "1. 切换到文件上传模式，显示「选择文件」按钮和只读文件名输入框\n2. 文件选择器仅接受.mp4格式\n3. 选择文件后，文件名正确显示在输入框中\n4. 格式、大小、时长校验通过，自动展示「正在处理视频...」loading条\n5. loading完成后显示「选择复刻视频时间段」按钮",
        "requirement": "文件上传-格式/大小/时长校验"
    },
    {
        "id": "TC-004",
        "module": "文件上传校验",
        "name": "文件上传-格式校验：非.mp4格式",
        "priority": "P0",
        "precondition": "复刻视频面板-文件上传模式",
        "steps": "1. 点击「选择文件」\n2. 选择一个非.mp4格式的视频文件（如.mov、.avi、.wmv）\n3. 选择一个非视频文件（如.jpg、.png）\n4. 观察错误提示",
        "expected": "1. 系统弹出提示「视频格式仅支持 .mp4，请重新选择」\n2. 文件名输入框内容清空\n3. 文件输入框重置\n4. 不触发后续下载/校验流程",
        "requirement": "视频格式仅支持.mp4"
    },
    {
        "id": "TC-005",
        "module": "文件上传校验",
        "name": "文件上传-大小校验：超过100MB",
        "priority": "P0",
        "precondition": "复刻视频面板-文件上传模式",
        "steps": "1. 选择一个.mp4格式但文件大小>100MB的视频\n2. 观察系统响应",
        "expected": "1. 系统弹出提示「文件大小不能超过 100MB，当前文件 XX MB，请重新选择」\n2. 文件名输入框内容清空\n3. 文件输入框重置\n4. 不触发后续流程",
        "requirement": "文件最大100MB"
    },
    {
        "id": "TC-006",
        "module": "文件上传校验",
        "name": "文件上传-大小边界：刚好100MB",
        "priority": "P1",
        "precondition": "复刻视频面板-文件上传模式",
        "steps": "1. 选择一个.mp4格式，文件大小恰好为 104857600 字节（100MB）的视频\n2. 观察校验结果",
        "expected": "1. 大小校验通过（100MB及以内均可接受）\n2. 文件被接受，进入后续处理流程\n3. 文件名正确显示",
        "requirement": "文件最大100MB-边界值"
    },
    {
        "id": "TC-007",
        "module": "文件上传校验",
        "name": "文件上传-时长校验：超过30秒",
        "priority": "P0",
        "precondition": "复刻视频面板-文件上传模式",
        "steps": "1. 选择一个.mp4格式，文件大小合规，但时长>30秒的视频\n2. 等待系统读取视频元数据\n3. 观察错误提示",
        "expected": "1. 系统读取视频元数据后，弹出提示「视频时长不能超过 30 秒，当前视频 XX 秒，请重新选择」\n2. 文件名输入框清空\n3. 文件输入框重置\n4. 不触发后续流程",
        "requirement": "视频最长30秒"
    },
    {
        "id": "TC-008",
        "module": "文件上传校验",
        "name": "文件上传-时长边界：刚好30秒",
        "priority": "P1",
        "precondition": "复刻视频面板-文件上传模式",
        "steps": "1. 选择一个.mp4格式，时长恰好30秒的视频\n2. 观察校验结果",
        "expected": "1. 时长校验通过（30秒及以内均可接受）\n2. 文件被接受，进入处理流程\n3. 视频时长正确记录（_cloneVideoDuration = 30秒）",
        "requirement": "视频最长30秒-边界值"
    },
    {
        "id": "TC-009",
        "module": "文件上传校验",
        "name": "文件上传-损坏或无效视频文件",
        "priority": "P2",
        "precondition": "复刻视频面板-文件上传模式",
        "steps": "1. 选择一个扩展名为.mp4但内容损坏的文件\n2. 等待系统尝试读取元数据\n3. 观察系统响应",
        "expected": "1. 系统弹出提示「无法读取视频信息，请确认文件是否为有效的 .mp4 视频」\n2. 文件名输入框清空\n3. 文件输入框重置\n4. 不触发后续流程",
        "requirement": "异常处理-无效视频文件"
    },
    {
        "id": "TC-010",
        "module": "文件上传校验",
        "name": "文件上传-取消选择文件",
        "priority": "P2",
        "precondition": "复刻视频面板-文件上传模式",
        "steps": "1. 点击「选择文件」\n2. 在文件选择器中点击「取消」不选择任何文件\n3. 观察面板状态",
        "expected": "1. 文件选择器关闭\n2. 面板状态不变，文件名仍显示「未选择文件」\n3. 不触发任何校验流程",
        "requirement": "文件上传-取消操作"
    },

    # ===== 模块3: 视频上传 - 链接上传校验 =====
    {
        "id": "TC-011",
        "module": "链接上传校验",
        "name": "链接上传方式-输入有效URL并确认",
        "priority": "P0",
        "precondition": "1. 复刻视频面板已打开\n2. 默认选中「链接」上传方式",
        "steps": "1. 在视频URL输入框中输入有效的.mp4视频链接\n2. 点击「确认」按钮\n3. 观察下载progress和校验结果",
        "expected": "1. 显示下载loading条，标题变为「正在下载视频...」\n2. 进度条从0%增长到100%\n3. 下载完成后显示「正在校验...」\n4. 校验通过后显示「校验通过」，loading条消失\n5. 显示「选择复刻视频时间段」按钮",
        "requirement": "链接上传-下载到本地并校验"
    },
    {
        "id": "TC-012",
        "module": "链接上传校验",
        "name": "链接上传-URL回车触发确认",
        "priority": "P1",
        "precondition": "复刻视频面板-链接上传模式",
        "steps": "1. 在URL输入框中输入视频链接\n2. 按Enter键\n3. 观察是否触发下载",
        "expected": "1. 按Enter键等价于点击「确认」按钮\n2. 触发下载流程",
        "requirement": "链接上传-交互便捷性"
    },
    {
        "id": "TC-013",
        "module": "链接上传校验",
        "name": "链接上传-格式校验：非.mp4链接被拒绝",
        "priority": "P0",
        "precondition": "复刻视频面板-链接上传模式",
        "steps": "1. 输入一个非.mp4格式的视频链接\n2. 点击「确认」\n3. 等待下载完成后模拟校验失败\n4. 观察错误提示",
        "expected": "1. 下载完成后进行格式校验\n2. 格式不匹配时给出明确提示\n3. loading条消失，不影响面板其他功能",
        "requirement": "链接上传-格式校验"
    },
    {
        "id": "TC-014",
        "module": "链接上传校验",
        "name": "链接上传-大小校验：超过100MB被拒绝",
        "priority": "P0",
        "precondition": "复刻视频面板-链接上传模式",
        "steps": "1. 输入一个指向大文件（>100MB）的视频链接\n2. 点击「确认」等待下载\n3. 校验阶段观察提示",
        "expected": "1. 校验时弹出提示「视频文件大小超过 100MB（XX MB），请更换视频」\n2. loading条消失\n3. 不进入时间段选择流程",
        "requirement": "链接上传-大小校验"
    },
    {
        "id": "TC-015",
        "module": "链接上传校验",
        "name": "链接上传-时长校验：超过30秒被拒绝",
        "priority": "P0",
        "precondition": "复刻视频面板-链接上传模式",
        "steps": "1. 输入一个时长>30秒的视频链接\n2. 点击「确认」等待下载和校验\n3. 观察校验阶段提示",
        "expected": "1. 校验时弹出提示「视频时长超过 30 秒（XX 秒），请更换视频」\n2. loading条消失\n3. 不进入时间段选择流程",
        "requirement": "链接上传-时长校验"
    },
    {
        "id": "TC-016",
        "module": "链接上传校验",
        "name": "链接上传-空URL点击确认",
        "priority": "P1",
        "precondition": "复刻视频面板-链接上传模式",
        "steps": "1. 不输入URL，直接点击「确认」按钮\n2. 观察系统反应",
        "expected": "1. URL为空时点击确认无响应（不触发下载）\n2. 无异常弹窗或崩溃",
        "requirement": "链接上传-空输入防护"
    },
    {
        "id": "TC-017",
        "module": "链接上传校验",
        "name": "链接上传-上传方式切换联动",
        "priority": "P1",
        "precondition": "复刻视频面板已打开",
        "steps": "1. 默认链接模式下确认UI显示\n2. 切换到「文件上传」模式\n3. 再切换回「链接」模式\n4. 观察各区域显示/隐藏状态",
        "expected": "1. 链接模式：显示URL输入框+确认按钮，隐藏文件选择区域\n2. 文件上传模式：显示文件选择区域，隐藏URL输入框\n3. 切换正确联动，无残留显示",
        "requirement": "上传方式切换"
    },

    # ===== 模块4: 时间段选择器 =====
    {
        "id": "TC-018",
        "module": "时间段选择器",
        "name": "打开时间段选择器-默认片段",
        "priority": "P0",
        "precondition": "1. 视频已上传/下载并校验通过\n2. 已显示「选择复刻视频时间段」按钮",
        "steps": "1. 点击「选择复刻视频时间段」按钮\n2. 观察时间段选择器弹窗\n3. 检查默认生成的片段",
        "expected": "1. 弹窗正确显示，标题为「选择复刻视频时间段」\n2. 显示视频播放区、视频帧预览条\n3. 默认生成1个时间段片段（默认7秒）\n4. 弹窗提示「每段最短 4 秒，最长 10 秒」",
        "requirement": "时间段选择器-打开"
    },
    {
        "id": "TC-019",
        "module": "时间段选择器",
        "name": "时间段-添加新片段",
        "priority": "P0",
        "precondition": "时间段选择器弹窗已打开",
        "steps": "1. 观察当前片段数量\n2. 点击「添加时间段」按钮\n3. 观察新增片段\n4. 再次点击添加2个",
        "expected": "1. 每次点击添加一个7秒默认片段\n2. 片段编号正确递增（片段1、片段2、片段3...）\n3. 新片段不影响已有片段的时间范围\n4. 删除按钮在所有片段上均显示（片段数>1时）",
        "requirement": "时间段选择器-添加片段"
    },
    {
        "id": "TC-020",
        "module": "时间段选择器",
        "name": "时间段-删除片段",
        "priority": "P1",
        "precondition": "时间段选择器中有多个片段（≥2个）",
        "steps": "1. 记录当前片段列表\n2. 点击某个片段的「删除」按钮\n3. 观察片段列表变化\n4. 当只剩1个片段时，观察删除按钮状态",
        "expected": "1. 被点击的片段从列表中移除\n2. 剩余片段编号自动重新排序（片段1、片段2...）\n3. 仅剩1个片段时，删除按钮隐藏（不可删除最后一个）",
        "requirement": "时间段选择器-删除片段"
    },
    {
        "id": "TC-021",
        "module": "时间段选择器",
        "name": "时间段-滑块调整时间范围",
        "priority": "P0",
        "precondition": "时间段选择器中有片段",
        "steps": "1. 拖动片段的起始滑块（ts-range-start）\n2. 拖动片段的结束滑块（ts-range-end）\n3. 观察时间显示更新\n4. 确认滑块范围限制在0到视频总时长之间",
        "expected": "1. 时间显示实时更新为「X s – Y s（Z 秒）」\n2. 滑块step为0.5秒\n3. 滑块范围限制在 0 到视频总时长\n4. 起始值不能超过结束值",
        "requirement": "时间段选择器-滑块调节"
    },
    {
        "id": "TC-022",
        "module": "时间段选择器",
        "name": "时间段-最短4秒限制校验",
        "priority": "P0",
        "precondition": "时间段选择器中至少有一个片段",
        "steps": "1. 将某个片段时间范围调整为不足4秒（如：0s–3.5s）\n2. 观察时间显示区域的警告提示\n3. 点击「确认选择」\n4. 观察是否弹出确认框",
        "expected": "1. 时间显示下方出现警告「⚠️ 片段时长不能少于 4 秒」\n2. 点击确认时弹出confirm提示「存在不符合 4-10 秒限制的时间段，确定继续？」\n3. 选择「取消」则不提交，选择「确定」则强制提交",
        "requirement": "每段最短4秒"
    },
    {
        "id": "TC-023",
        "module": "时间段选择器",
        "name": "时间段-最长10秒限制校验",
        "priority": "P0",
        "precondition": "时间段选择器中至少有一个片段",
        "steps": "1. 将某个片段时间范围调整为超过10秒（如：0s–12s）\n2. 观察警告提示\n3. 点击确认选择",
        "expected": "1. 出现警告「⚠️ 片段时长不能超过 10 秒」\n2. 确认时有confirm提示\n3. 可选择取消或强制提交",
        "requirement": "每段最长10秒"
    },
    {
        "id": "TC-024",
        "module": "时间段选择器",
        "name": "时间段-边界值：刚好4秒和刚好10秒",
        "priority": "P1",
        "precondition": "时间段选择器已打开",
        "steps": "1. 将片段1调整为刚好4秒（如：0s–4s）\n2. 添加片段2，调整为刚好10秒（如：4s–14s）\n3. 确认选择",
        "expected": "1. 4秒和10秒均无警告提示\n2. 确认选择时不弹出confirm警告\n3. 时间段正常提交到面板",
        "requirement": "4秒和10秒边界值"
    },
    {
        "id": "TC-025",
        "module": "时间段选择器",
        "name": "自动选择-按总时长均分",
        "priority": "P1",
        "precondition": "时间段选择器已打开，视频总时长≥14秒",
        "steps": "1. 点击「自动选择」按钮\n2. 观察生成的片段数量和分布\n3. 验证每个片段时长在4-10秒范围内",
        "expected": "1. 按视频总时长余10秒后，以7秒为基准均分\n2. 所有片段时长在4-10秒之间\n3. 片段时间连续覆盖，无重叠无断档\n4. 如视频时长不足14秒，弹出提示「视频时长不足，无法自动选择（至少需要 14 秒）」",
        "requirement": "自动选择-智能均分"
    },
    {
        "id": "TC-026",
        "module": "时间段选择器",
        "name": "确认时间段选择-返回面板展示",
        "priority": "P0",
        "precondition": "时间段选择器弹窗已配置好合法时间段",
        "steps": "1. 确认所有片段均在4-10秒内\n2. 点击「确认选择」按钮\n3. 观察弹窗关闭和面板变化",
        "expected": "1. 弹窗关闭\n2. 复刻面板显示「已选时间段」区域\n3. 每个片段显示：编号、时间范围、时长、预计消耗5积分\n4. 提供提交按钮",
        "requirement": "确认选择-返回面板"
    },
    {
        "id": "TC-027",
        "module": "时间段选择器",
        "name": "时间段选择器-取消操作",
        "priority": "P1",
        "precondition": "时间段选择器已打开",
        "steps": "1. 配置好时间段\n2. 点击「取消」按钮\n3. 观察弹窗和面板状态",
        "expected": "1. 弹窗关闭\n2. 时间段选择不保存\n3. 下次打开选择器时重新初始化（默认1个7秒片段）",
        "requirement": "时间段选择器-取消"
    },

    # ===== 模块5: 任务提交与列表 =====
    {
        "id": "TC-028",
        "module": "任务提交与列表",
        "name": "复刻视频任务-完整提交流程",
        "priority": "P0",
        "precondition": "1. 视频名称已填写\n2. 复刻内容描述已填写\n3. 时间段已选择（全部4-10秒内）",
        "steps": "1. 确认所有必填项已填写\n2. 点击弹窗底部「提交」按钮\n3. 观察弹窗关闭和左侧任务列表",
        "expected": "1. 弹窗关闭\n2. 左侧任务列表新增一条复刻视频任务\n3. 任务卡片显示：橙色「复刻中」状态角标、视频名称、片段数量、总时长、橙色「复刻视频」类型标签\n4. 任务卡片展开显示子任务列表（每个片段一行）\n5. 复刻视频面板表单重置",
        "requirement": "任务提交-完整流程"
    },
    {
        "id": "TC-029",
        "module": "任务提交与列表",
        "name": "提交校验-视频名称为空",
        "priority": "P1",
        "precondition": "复刻视频面板中时间段已选择",
        "steps": "1. 不填写视频名称\n2. 填写复刻内容描述\n3. 点击「提交」\n4. 观察提示",
        "expected": "1. 弹出提示「请输入视频名称」\n2. 弹窗不关闭\n3. 提交被阻止",
        "requirement": "必填校验-视频名称"
    },
    {
        "id": "TC-030",
        "module": "任务提交与列表",
        "name": "提交校验-复刻内容描述为空",
        "priority": "P1",
        "precondition": "复刻视频面板中时间段已选择",
        "steps": "1. 填写视频名称\n2. 不填写复刻内容描述\n3. 点击「提交」",
        "expected": "1. 弹出提示「请输入复刻内容描述」\n2. 弹窗不关闭\n3. 提交被阻止",
        "requirement": "必填校验-复刻内容描述"
    },
    {
        "id": "TC-031",
        "module": "任务提交与列表",
        "name": "提交校验-未选择时间段",
        "priority": "P0",
        "precondition": "复刻视频面板中视频名称和描述已填写，但未选择时间段",
        "steps": "1. 填写视频名称和描述\n2. 不点击「选择复刻视频时间段」\n3. 直接点击「提交」",
        "expected": "1. 弹出提示「请先选择复刻视频时间段」\n2. 弹窗不关闭\n3. 提交被阻止",
        "requirement": "必填校验-时间段"
    },
    {
        "id": "TC-032",
        "module": "任务提交与列表",
        "name": "提交校验-时间段存在非法时长（4-10秒外）",
        "priority": "P1",
        "precondition": "已选择时间段，但存在<4秒或>10秒的片段",
        "steps": "1. 选择时间段时强制提交了不合规片段\n2. 点击「提交」\n3. 观察提示",
        "expected": "1. 弹出confirm提示「存在不符合 4-10 秒限制的时间段，确定提交？」\n2. 选择「取消」不提交\n3. 选择「确定」允许强制提交",
        "requirement": "提交时二次确认"
    },
    {
        "id": "TC-033",
        "module": "任务提交与列表",
        "name": "任务列表-复刻视频任务卡片展示",
        "priority": "P0",
        "precondition": "已成功创建复刻视频任务",
        "steps": "1. 查看左侧任务列表\n2. 观察新创建的任务卡片\n3. 点击卡片头部选中\n4. 点击展开/折叠按钮",
        "expected": "1. 任务卡片显示复刻图标（橙色播放按钮）\n2. 卡片hint显示「复刻视频 · N个片段 · XX秒」\n3. 类型标签为橙色「复刻视频」\n4. 状态角标显示「复刻中」\n5. 卡片选中时右侧栏切换为复刻视频模式（Tab：视频方案、时间段）\n6. 子任务列表正确展开/折叠",
        "requirement": "任务列表-卡片展示"
    },
    {
        "id": "TC-034",
        "module": "任务提交与列表",
        "name": "右侧面板-复刻视频方案展示",
        "priority": "P1",
        "precondition": "已选中复刻视频任务卡片",
        "steps": "1. 选中复刻视频任务\n2. 观察右侧面板\n3. 查看「视频方案」Tab内容\n4. 切换到「时间段」Tab查看",
        "expected": "1. 右侧面板切换到复刻视频模式\n2. Tab显示：视频方案 → 时间段\n3. 视频方案卡片展示：视频名称、片段数量、总时长\n4. 时间段Tab展示每个片段的卡片：编号、片段名称、时间范围、时长\n5. 底部显示「生成视频」和「下载视频」按钮",
        "requirement": "右侧面板-方案展示"
    },

    # ===== 模块6: 积分消耗 =====
    {
        "id": "TC-035",
        "module": "积分消耗",
        "name": "积分计算-按片段个数扣取",
        "priority": "P0",
        "precondition": "复刻视频面板中已选择N个时间段",
        "steps": "1. 选择1个片段，查看积分提示\n2. 选择3个片段，查看积分提示\n3. 选择5个片段，查看积分提示",
        "expected": "1. 每个片段显示「预计消耗 5 积分」\n2. 1个片段 → 5积分\n3. 3个片段 → 15积分\n4. 5个片段 → 25积分\n5. 积分提示在每个片段行中正确显示",
        "requirement": "每个片段消耗5积分"
    },
    {
        "id": "TC-036",
        "module": "积分消耗",
        "name": "积分计算-自动选择片段数×5",
        "priority": "P1",
        "precondition": "使用自动选择功能分配了N个片段",
        "steps": "1. 对一段视频使用「自动选择」\n2. 观察生成的片段数\n3. 累计所有片段的积分消耗",
        "expected": "1. 每个片段均显示5积分\n2. 总积分 = 片段数 × 5\n3. 积分提示与手动添加片段一致",
        "requirement": "积分计算-自动选择"
    },
    {
        "id": "TC-037",
        "module": "积分消耗",
        "name": "积分消耗-用户积分不足场景",
        "priority": "P2",
        "precondition": "用户账户积分余额 < 所需消耗积分",
        "steps": "1. 选择N个片段（所需积分 > 余额）\n2. 点击「提交」\n3. 观察系统提示",
        "expected": "1. 提交时弹出积分不足提示\n2. 明确告知所需积分和当前余额\n3. 引导用户充值或减少片段数量",
        "requirement": "积分不足-异常处理"
    },
    {
        "id": "TC-038",
        "module": "积分消耗",
        "name": "积分消耗-删除片段后积分重新计算",
        "priority": "P1",
        "precondition": "时间段选择器中已添加多个片段",
        "steps": "1. 添加3个片段（预计15积分）\n2. 删除1个片段\n3. 确认选择后查看积分\n4. 观察面板展示",
        "expected": "1. 确认选择后面板仅显示2个片段\n2. 每个片段均显示5积分\n3. 总积分从15变为10",
        "requirement": "积分重算-删除片段"
    },
]

wb = Workbook()

# ===== Sheet 1: 冒烟测试用例 =====
ws = wb.active
ws.title = "复刻视频-冒烟测试用例"

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

header_fill = PatternFill(start_color="E65100", end_color="E65100", fill_type="solid")
header_font = Font(bold=True, size=11, color="FFFFFF", name="微软雅黑")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

priority_colors = {
    "P0": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "P1": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "P2": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
}

priority_fonts = {
    "P0": Font(bold=True, size=10, color="9C0006", name="微软雅黑"),
    "P1": Font(bold=True, size=10, color="9C6500", name="微软雅黑"),
    "P2": Font(bold=True, size=10, color="006100", name="微软雅黑"),
}

module_colors = {
    "任务创建-面板切换": "FFF3E0",
    "文件上传校验": "FFF8E1",
    "链接上传校验": "FFFDE7",
    "时间段选择器": "F3E5F5",
    "任务提交与列表": "E8F5E9",
    "积分消耗": "E3F2FD",
}

headers = ["用例ID", "模块", "用例名称", "优先级", "前置条件", "测试步骤", "预期结果", "关联需求"]

for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment
    cell.border = thin_border

for row_idx, case in enumerate(test_cases, 2):
    data = [
        case["id"],
        case["module"],
        case["name"],
        case["priority"],
        case["precondition"],
        case["steps"],
        case["expected"],
        case["requirement"],
    ]
    module_fill = PatternFill(
        start_color=module_colors.get(case["module"], "FFFFFF"),
        end_color=module_colors.get(case["module"], "FFFFFF"),
        fill_type="solid"
    )
    for col_idx, value in enumerate(data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=10, name="微软雅黑")

        if col_idx == 1:
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif col_idx == 2:
            cell.fill = module_fill
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif col_idx == 4:
            cell.fill = priority_colors.get(value, None)
            cell.font = priority_fonts.get(value, Font(size=10, name="微软雅黑"))
            cell.alignment = Alignment(horizontal="center", vertical="top")

column_widths = {
    1: 10, 2: 16, 3: 32, 4: 8, 5: 32, 6: 48, 7: 42, 8: 22
}
for col_idx, width in column_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 30
for row_idx in range(2, len(test_cases) + 2):
    ws.row_dimensions[row_idx].height = 85

ws.auto_filter.ref = f"A1:H{len(test_cases) + 1}"
ws.freeze_panes = "A2"

# ===== Sheet 2: 测试用例统计 =====
ws2 = wb.create_sheet(title="测试用例统计")

stat_header_fill = PatternFill(start_color="E65100", end_color="E65100", fill_type="solid")
stat_header_font = Font(bold=True, size=12, color="FFFFFF", name="微软雅黑")
section_font = Font(bold=True, size=11, color="E65100", name="微软雅黑")
section_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")

ws2.column_dimensions['A'].width = 28
ws2.column_dimensions['B'].width = 12
ws2.column_dimensions['C'].width = 12
ws2.column_dimensions['D'].width = 12
ws2.column_dimensions['E'].width = 12

ws2.merge_cells('A1:E1')
title_cell = ws2.cell(row=1, column=1, value="复刻视频 - 新建复刻视频流程 - 冒烟测试用例统计")
title_cell.font = Font(bold=True, size=14, color="BF360C", name="微软雅黑")
title_cell.alignment = Alignment(horizontal="center", vertical="center")
title_cell.fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
ws2.row_dimensions[1].height = 36

row = 3
ws2.cell(row=row, column=1, value="一、模块分布统计").font = section_font
for c in range(1, 6):
    ws2.cell(row=row, column=c).fill = section_fill
    ws2.cell(row=row, column=c).border = thin_border

row += 1
stat_headers = ["模块名称", "P0数量", "P1数量", "P2数量", "小计"]
for col_idx, h in enumerate(stat_headers, 1):
    cell = ws2.cell(row=row, column=col_idx, value=h)
    cell.fill = stat_header_fill
    cell.font = stat_header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border
ws2.row_dimensions[row].height = 26

module_order = ["任务创建-面板切换", "文件上传校验", "链接上传校验",
                "时间段选择器", "任务提交与列表", "积分消耗"]
module_data = {}
for tc in test_cases:
    m = tc["module"]
    if m not in module_data:
        module_data[m] = {"P0": 0, "P1": 0, "P2": 0}
    module_data[m][tc["priority"]] += 1

total_p0 = total_p1 = total_p2 = 0
row += 1
for mod in module_order:
    if mod in module_data:
        d = module_data[mod]
        vals = [mod, d["P0"], d["P1"], d["P2"], d["P0"] + d["P1"] + d["P2"]]
        total_p0 += d["P0"]
        total_p1 += d["P1"]
        total_p2 += d["P2"]
        for col_idx, v in enumerate(vals, 1):
            cell = ws2.cell(row=row, column=col_idx, value=v)
            cell.border = thin_border
            cell.font = Font(size=10, name="微软雅黑")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[row].height = 24
        row += 1

row_data = ["合计", total_p0, total_p1, total_p2, total_p0 + total_p1 + total_p2]
for col_idx, v in enumerate(row_data, 1):
    cell = ws2.cell(row=row, column=col_idx, value=v)
    cell.border = thin_border
    cell.font = Font(bold=True, size=11, name="微软雅黑")
    cell.fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[row].height = 26

row += 2
ws2.cell(row=row, column=1, value="二、优先级分布统计").font = section_font
for c in range(1, 6):
    ws2.cell(row=row, column=c).fill = section_fill
    ws2.cell(row=row, column=c).border = thin_border

row += 1
pri_headers = ["优先级", "用例数量", "占比", "说明", ""]
for col_idx, h in enumerate(pri_headers, 1):
    cell = ws2.cell(row=row, column=col_idx, value=h)
    cell.fill = stat_header_fill
    cell.font = stat_header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border
ws2.row_dimensions[row].height = 26

pri_data = [
    ("P0 - 核心流程", total_p0, f"{total_p0/len(test_cases)*100:.1f}%", "核心业务流程，必须通过", "FFC7CE"),
    ("P1 - 重要功能", total_p1, f"{total_p1/len(test_cases)*100:.1f}%", "重要功能/校验/边界，应通过", "FFEB9C"),
    ("P2 - 异常/体验", total_p2, f"{total_p2/len(test_cases)*100:.1f}%", "异常处理/用户体验，建议通过", "C6EFCE"),
]

row += 1
for p_name, p_count, p_ratio, p_desc, p_color in pri_data:
    vals = [p_name, p_count, p_ratio, p_desc]
    for col_idx, v in enumerate(vals, 1):
        cell = ws2.cell(row=row, column=col_idx, value=v)
        cell.border = thin_border
        cell.font = Font(size=10, name="微软雅黑")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill(start_color=p_color, end_color=p_color, fill_type="solid")
    ws2.row_dimensions[row].height = 28
    row += 1

row += 1
ws2.merge_cells(f'A{row}:E{row}')
cell = ws2.cell(row=row, column=1, value=f"总计：{len(test_cases)} 条冒烟测试用例")
cell.font = Font(bold=True, size=12, color="BF360C", name="微软雅黑")
cell.alignment = Alignment(horizontal="center", vertical="center")
cell.border = thin_border
ws2.row_dimensions[row].height = 30

row += 2
ws2.cell(row=row, column=1, value="三、测试覆盖说明").font = section_font
for c in range(1, 6):
    ws2.cell(row=row, column=c).fill = section_fill
    ws2.cell(row=row, column=c).border = thin_border

row += 1
info_items = [
    "测试范围：AI视频工具3.1 - 新建复刻视频流程（冒烟测试）",
    "需求版本：AI视频工具3.1 - 复刻视频功能（一期新增）",
    "设计日期：2026-06-01",
    "覆盖功能模块：6个",
    "主要流程覆盖：全覆盖（面板切换 → 视频上传校验 → 时间段选择 → 任务提交 → 积分消耗）",
    "核心校验规则：",
    "  1. 视频格式仅支持 .mp4",
    "  2. 文件上传 ≤ 100MB，链接下载后校验 ≤ 100MB",
    "  3. 视频时长 ≤ 30秒",
    "  4. 时间段每段最短 4秒，最长 10秒",
    "  5. 积分按片段个数扣取，每个片段 5 积分",
    "未覆盖项：",
    "  - 视频实际播放与帧截取功能（需接入真实视频编解码）",
    "  - 链接下载真实网络异常处理（需接入实际下载模块）",
    "  - 视频生成和下载实际功能（非复刻流程特有）",
]
for item in info_items:
    cell = ws2.cell(row=row, column=1, value=item)
    ws2.merge_cells(f'A{row}:E{row}')
    cell.font = Font(size=10, name="微软雅黑")
    cell.alignment = Alignment(vertical="center")
    cell.border = thin_border
    ws2.row_dimensions[row].height = 20
    row += 1

output_path = "/Users/lanwang/Documents/trae_projects/new World/复刻视频_新建流程_冒烟测试用例.xlsx"
wb.save(output_path)
print(f"复刻视频冒烟测试用例Excel已生成：{output_path}")
print(f"   共 {len(test_cases)} 条测试用例")
print(f"   P0: {total_p0} 条 | P1: {total_p1} 条 | P2: {total_p2} 条")

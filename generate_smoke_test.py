# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter
from collections import Counter

test_cases = [
    # ===== 模块1: 任务创建与管理 =====
    {
        "id": "TC-001",
        "module": "任务创建与管理",
        "name": "AI自动生成脚本任务-完整创建流程",
        "priority": "P0",
        "precondition": "1. 用户已登录系统\n2. 处于AI视频创作平台首页",
        "steps": "1. 点击左侧栏「+ 新建任务」按钮\n2. 在弹窗中点击「AI自动生成脚本」TAB\n3. 输入商品名称（如：玻尿酸保湿面膜 50片装）\n4. 选择商品品类\n5. 选择商品信息来源为「商品URL」\n6. 输入有效的商品URL\n7. 输入生成视频长度（如：45秒）\n8. 点击「提交」按钮",
        "expected": "1. 弹窗标题切换为「AI自动生成脚本任务」\n2. AI脚本提示文案正确展示\n3. 各输入项正常可编辑\n4. 提交后弹窗关闭，任务列表中新增一条AI自动生成脚本任务\n5. 任务卡片显示正确的商品名称和任务类型标签",
        "requirement": "需求1-步骤1"
    },
    {
        "id": "TC-002",
        "module": "任务创建与管理",
        "name": "视频解析任务-完整创建流程",
        "priority": "P0",
        "precondition": "1. 用户已登录系统\n2. 处于AI视频创作平台首页",
        "steps": "1. 点击左侧栏「+ 新建任务」按钮\n2. 确认弹窗默认选中「视频解析任务」TAB\n3. 选择图片生成模型（如：豆包）\n4. 输入视频长度（如：30秒）\n5. 选择商品上传方式为「链接」\n6. 选择商品品类\n7. 输入商品URL\n8. 选择视频上传方式为「链接」\n9. 输入视频URL\n10. 点击「提交」按钮",
        "expected": "1. 弹窗标题显示「新增视频解析任务」\n2. 所有表单字段正常可编辑\n3. 提交后弹窗关闭，任务列表新增一条视频解析任务\n4. 任务卡片显示正确信息",
        "requirement": "需求1-步骤1"
    },
    {
        "id": "TC-003",
        "module": "任务创建与管理",
        "name": "任务创建弹窗-任务类型切换",
        "priority": "P0",
        "precondition": "用户已打开新建任务弹窗",
        "steps": "1. 确认弹窗默认显示「视频解析任务」表单\n2. 点击「AI自动生成脚本」TAB\n3. 观察表单内容是否切换\n4. 再次点击「视频解析任务」TAB\n5. 观察表单内容是否切回",
        "expected": "1. 默认显示视频解析任务表单\n2. 切换到AI表单后，显示AI脚本提示、商品名称、品类、信息来源等字段\n3. 切回后还原为视频解析任务表单\n4. TAB切换流畅，UI高亮状态正确",
        "requirement": "需求1-步骤1"
    },
    {
        "id": "TC-004",
        "module": "任务创建与管理",
        "name": "AI任务-商品信息来源切换（URL/文件上传）",
        "priority": "P1",
        "precondition": "打开AI自动生成脚本任务弹窗",
        "steps": "1. 确认默认选中「商品URL」\n2. 观察URL输入框正常显示\n3. 点击「商品图片上传」单选\n4. 观察文件上传区域是否显示\n5. 点击「商品URL」切回",
        "expected": "1. 切换时对应输入区域正确显示/隐藏\n2. URL模式下显示输入框，文件模式下显示上传区域\n3. 切换后的状态正确保持",
        "requirement": "需求1-步骤1a"
    },
    {
        "id": "TC-005",
        "module": "任务创建与管理",
        "name": "AI任务-必填字段校验",
        "priority": "P1",
        "precondition": "打开AI自动生成脚本任务弹窗",
        "steps": "1. 不填写商品名称，直接点击「提交」\n2. 填写商品名称，不选择品类，点击「提交」\n3. 填写商品名称和品类，不填写视频长度，点击「提交」",
        "expected": "1. 弹出提示「请输入商品名称」\n2. 弹出提示「请选择商品品类」\n3. 弹出提示「请输入生成视频长度」\n4. 每次校验通过后方可提交",
        "requirement": "需求1-步骤1a"
    },
    {
        "id": "TC-006",
        "module": "任务创建与管理",
        "name": "任务列表-任务卡片展示",
        "priority": "P0",
        "precondition": "系统中存在多个已创建的任务",
        "steps": "1. 查看左侧栏任务列表\n2. 观察每个任务卡片的展示内容\n3. 查看任务类型标签\n4. 查看任务状态角标",
        "expected": "1. 任务卡片按创建时间倒序排列\n2. 每个卡片显示：缩略图、任务名称、提示信息、类型标签\n3. AI自动生成脚本任务显示紫色标签\n4. 视频解析任务显示蓝色标签\n5. 任务状态角标正常显示（成功/失败）",
        "requirement": "需求1"
    },
    {
        "id": "TC-007",
        "module": "任务创建与管理",
        "name": "任务列表-搜索功能",
        "priority": "P1",
        "precondition": "任务列表中有多个任务",
        "steps": "1. 在搜索输入框中输入任务名称关键词\n2. 观察列表过滤结果\n3. 点击清除按钮\n4. 观察列表是否恢复",
        "expected": "1. 输入关键词后，列表仅展示匹配的任务\n2. 清除搜索内容后，列表恢复全部展示\n3. 搜索输入框的清除按钮在有内容时显示",
        "requirement": "HTML原型"
    },
    {
        "id": "TC-008",
        "module": "任务创建与管理",
        "name": "任务列表-分类标签筛选",
        "priority": "P1",
        "precondition": "任务列表中有不同类型的任务",
        "steps": "1. 点击分类标签（如「AI自动生成脚本」）\n2. 观察列表过滤结果\n3. 再次点击该标签取消筛选\n4. 观察列表恢复情况",
        "expected": "1. 点击后只显示对应类型的任务\n2. 选中标签高亮为紫色\n3. 取消筛选后列表恢复全部\n4. 标签样式恢复正常",
        "requirement": "HTML原型"
    },

    # ===== 模块2: AI自动生成脚本-商品信息处理 =====
    {
        "id": "TC-009",
        "module": "商品信息处理",
        "name": "商品信息输入-URL方式解析",
        "priority": "P0",
        "precondition": "已创建AI自动生成脚本任务，进入方案确认页",
        "steps": "1. 在新建任务弹窗中输入有效商品URL\n2. 提交任务\n3. 系统自动爬取商品信息\n4. 观察右侧栏商品卖点区域",
        "expected": "1. 系统成功解析商品链接\n2. 获取商品详情中的图片链接\n3. 获取商品名称\n4. 解析结果在右侧商品卖点区域正确展示",
        "requirement": "需求1-步骤2"
    },
    {
        "id": "TC-010",
        "module": "商品信息处理",
        "name": "商品信息解析-数据完整性",
        "priority": "P0",
        "precondition": "已创建AI任务并输入了有效商品URL",
        "steps": "1. 提交任务后等待系统解析\n2. 检查解析出的商品名称\n3. 检查解析出的商品图片\n4. 检查右侧商品卖点Tab内容",
        "expected": "1. 商品名称完整准确\n2. 商品图片正常加载展示\n3. 卖点信息包含：产品名称、品类、属性、痛点、核心卖点、差异化卖点、目标人群\n4. 卖点内容符合合规检查要求",
        "requirement": "需求1-步骤3"
    },
    {
        "id": "TC-011",
        "module": "商品信息处理",
        "name": "卖点提炼-内容合规校验",
        "priority": "P0",
        "precondition": "商品信息已解析成功",
        "steps": "1. 查看卖点提炼结果\n2. 检查是否包含违禁词\n3. 检查合规检查字段（compliance_check）\n4. 验证差异化卖点数量为3个",
        "expected": "1. 不包含千川极限词（最/第一/顶级/唯一等）\n2. 不包含医疗功效词（治疗/治愈/根治等）\n3. 不包含虚假承诺词\n4. 核心卖点≤25字\n5. 差异化卖点描述≤30字\n6. 痛点描述≤30字\n7. 输出严格为JSON格式",
        "requirement": "需求1-步骤3"
    },

    # ===== 模块3: AI自动生成脚本-方案生成 =====
    {
        "id": "TC-012",
        "module": "方案生成",
        "name": "剧情脚本生成-内容完整性",
        "priority": "P0",
        "precondition": "1. AI任务已创建\n2. 商品信息已解析\n3. 卖点已提炼",
        "steps": "1. 系统自动调用剧情生成提示词\n2. 观察右侧「剧情脚本」Tab\n3. 查看每个镜头的卡片内容\n4. 检查镜头数量和时间分布",
        "expected": "1. 剧情脚本正确生成\n2. 每个镜头卡片包含：镜号、时间段、动作/台词、景别、机位/运动\n3. 镜头时间覆盖总时长，时间连续无断档\n4. 内容与卖点JSON、商品图片相关联",
        "requirement": "需求1-步骤5"
    },
    {
        "id": "TC-013",
        "module": "方案生成",
        "name": "视觉风格人物设定",
        "priority": "P0",
        "precondition": "1. 剧情脚本已生成\n2. 卖点已提炼",
        "steps": "1. 查看右侧「视觉风格」Tab\n2. 检查视觉风格字段内容\n3. 确认人物形象描述与商品风格匹配",
        "expected": "1. 视觉风格信息正确展示\n2. 包含风格关键词、色调、人物形象描述\n3. 视觉风格与商品品类、卖点匹配",
        "requirement": "需求1-步骤6"
    },
    {
        "id": "TC-014",
        "module": "方案生成",
        "name": "商品四宫格图片生成",
        "priority": "P1",
        "precondition": "1. 商品信息已解析\n2. 卖点已提炼",
        "steps": "1. 系统调用商品信息生图提示词\n2. 调用img2生成商品四宫格图片\n3. 查看右侧「素材图片」Tab中的产品图",
        "expected": "1. 商品四宫格图片成功生成\n2. 图片清晰，与商品详情匹配\n3. 四宫格图片展示在素材图片区域的产品图位置",
        "requirement": "需求1-步骤4"
    },
    {
        "id": "TC-015",
        "module": "方案生成",
        "name": "人物九宫格图片生成",
        "priority": "P1",
        "precondition": "1. 视觉风格人物设定已生成\n2. 人物信息生图提示词已配置",
        "steps": "1. 系统调用人物信息生图提示词\n2. 调用img2生成人物九宫格图片\n3. 查看右侧「素材图片」Tab中的人物图",
        "expected": "1. 人物九宫格图片成功生成\n2. 人物形象与视觉风格描述一致\n3. 图片展示在素材图片区域的人物图位置",
        "requirement": "需求1-步骤7"
    },

    # ===== 模块4: AI自动生成脚本-视频素材图片 =====
    {
        "id": "TC-016",
        "module": "视频素材图片生成",
        "name": "视频策划全览图生成",
        "priority": "P0",
        "precondition": "1. 卖点、剧情、视觉风格、人物九宫格、商品四宫格均已生成\n2. 右侧已切换到「视频素材图片」Tab",
        "steps": "1. 点击「生成视频素材图片」按钮\n2. 等待视频策划全览图生成\n3. 观察全览图区域的展示状态变化\n4. 点击「预览」按钮",
        "expected": "1. 按钮变为「生成中...」状态，带旋转动画\n2. 视频素材图片模块自动展开\n3. 视频策划全览图从「未生成」→「生成中...」→「已生成」\n4. 生成完成后显示：重新生成、替换、预览按钮\n5. 点击预览可放大查看图片",
        "requirement": "需求1-步骤8"
    },
    {
        "id": "TC-017",
        "module": "视频素材图片生成",
        "name": "分镜图生成",
        "priority": "P0",
        "precondition": "已触发视频素材图片生成流程",
        "steps": "1. 在视频策划全览图生成完成后\n2. 等待分镜图自动生成\n3. 观察分镜图区域状态\n4. 点击分镜图预览",
        "expected": "1. 分镜图在视频策划全览图之后自动生成\n2. 「未生成」→「生成中...」→「已生成」状态正确流转\n3. 生成完成后显示：重新生成、替换、预览按钮\n4. 分镜图与剧情脚本、视觉风格匹配",
        "requirement": "需求1-步骤9"
    },
    {
        "id": "TC-018",
        "module": "视频素材图片生成",
        "name": "视频素材图片生成-触发与防重复",
        "priority": "P0",
        "precondition": "1. 处于AI任务方案确认页\n2. 素材图片尚未生成",
        "steps": "1. 点击「生成视频素材图片」按钮\n2. 观察按钮状态变化\n3. 在生成过程中再次尝试点击按钮\n4. 等待生成完成后观察按钮状态",
        "expected": "1. 首次点击正常触发生成\n2. 生成中按钮禁用，显示loading状态\n3. 生成中再次点击无响应\n4. 生成完成后按钮变为「✓ 已生成」\n5. 外部生成按钮隐藏，底部显示「生成视频」和「下载视频」按钮",
        "requirement": "需求1-步骤8-9"
    },
    {
        "id": "TC-019",
        "module": "视频素材图片生成",
        "name": "重新生成视频素材图片",
        "priority": "P1",
        "precondition": "视频素材图片（全览图+分镜图）已生成",
        "steps": "1. 点击视频策划全览图的「重新生成」按钮\n2. 等待重新生成完成\n3. 点击分镜图的「重新生成」按钮\n4. 验证新生成的图片是否覆盖旧图",
        "expected": "1. 单个图片可以独立重新生成\n2. 重新生成时有loading状态\n3. 新图覆盖旧图展示\n4. 另一个图片不受影响",
        "requirement": "需求1-非顺序步骤2"
    },

    # ===== 模块5: 方案内容编辑与修改 =====
    {
        "id": "TC-020",
        "module": "方案内容编辑",
        "name": "商品卖点-内容编辑保存",
        "priority": "P1",
        "precondition": "商品卖点已生成并展示在右侧栏",
        "steps": "1. 在商品卖点区域点击编辑\n2. 修改核心卖点内容\n3. 点击「保存」\n4. 观察卖点内容是否更新",
        "expected": "1. 编辑态正确切换，显示输入框\n2. 修改保存后内容更新\n3. 保存成功后弹出「可重新生成视频素材图片」提示",
        "requirement": "需求1-非顺序步骤1"
    },
    {
        "id": "TC-021",
        "module": "方案内容编辑",
        "name": "剧情脚本-单镜头编辑保存",
        "priority": "P1",
        "precondition": "剧情脚本已生成，多个镜头卡片展示中",
        "steps": "1. 点击某个镜头卡片的编辑按钮\n2. 修改「动作/台词」字段\n3. 修改「景别」字段\n4. 点击「保存」\n5. 观察卡片内容是否更新",
        "expected": "1. 编辑态切换：卡片body隐藏，编辑区域显示\n2. 编辑区域包含：时间段、动作/台词、景别、机位/运动\n3. 保存后卡片展示更新后内容\n4. 弹窗提示可重新生成视频素材图片",
        "requirement": "需求1-非顺序步骤1"
    },
    {
        "id": "TC-022",
        "module": "方案内容编辑",
        "name": "视觉风格-编辑保存",
        "priority": "P1",
        "precondition": "视觉风格已生成展示",
        "steps": "1. 在视觉风格区域点击编辑\n2. 修改风格关键词\n3. 修改色调\n4. 点击「保存」",
        "expected": "1. 编辑态展示输入框，预填当前值\n2. 保存后内容更新\n3. 提示可重新生成视频素材图片",
        "requirement": "需求1-非顺序步骤1"
    },
    {
        "id": "TC-023",
        "module": "方案内容编辑",
        "name": "分镜脚本-镜头编辑保存",
        "priority": "P1",
        "precondition": "分镜脚本已生成（视频解析任务场景）",
        "steps": "1. 点击某个分镜卡片\n2. 在编辑区域修改运镜方式\n3. 修改视觉描述\n4. 修改台词\n5. 点击「保存」",
        "expected": "1. 分镜卡片选中高亮\n2. 编辑区域可修改运镜、视觉描述、通感、台词\n3. 保存后卡片视图更新\n4. 支持取消编辑恢复原值",
        "requirement": "需求1-步骤9-3"
    },
    {
        "id": "TC-024",
        "module": "方案内容编辑",
        "name": "素材图片-替换功能",
        "priority": "P1",
        "precondition": "素材图片已生成展示在网格中",
        "steps": "1. 在素材图片区域找到某张图片\n2. 点击图片上的「替换」按钮\n3. 选择新的图片文件\n4. 确认替换",
        "expected": "1. 点击替换后弹出文件选择器\n2. 选择文件后提示替换成功\n3. 图片区域显示新替换的图片\n4. 其他图片不受影响",
        "requirement": "需求1-非顺序步骤2"
    },
    {
        "id": "TC-025",
        "module": "方案内容编辑",
        "name": "修改后重新生成提醒",
        "priority": "P0",
        "precondition": "视频素材图片已生成",
        "steps": "1. 修改视觉风格内容并保存\n2. 观察右侧「视频素材图片」区域状态\n3. 修改剧情脚本内容并保存\n4. 观察状态变化",
        "expected": "1. 修改保存后，视频素材图片区域提示需重新生成\n2. 状态标签变为「未生成」或相应提示\n3. 「重新生成」按钮可用\n4. 各Agent修改后联动提示一致",
        "requirement": "需求1-非顺序步骤1"
    },

    # ===== 模块6: 视频生成与导出 =====
    {
        "id": "TC-026",
        "module": "视频生成与导出",
        "name": "生成视频-功能触发",
        "priority": "P0",
        "precondition": "1. AI任务方案确认页\n2. 视频素材图片已全部生成\n3. 底部显示「生成视频」按钮",
        "steps": "1. 确认素材图片全部已生成\n2. 点击「生成视频」按钮\n3. 等待视频生成",
        "expected": "1. 按钮正常可点击\n2. 点击后开始生成视频\n3. 生成过程有loading状态\n4. 视频生成完毕后可预览",
        "requirement": "需求1"
    },
    {
        "id": "TC-027",
        "module": "视频生成与导出",
        "name": "下载视频-功能触发",
        "priority": "P1",
        "precondition": "视频已生成完成",
        "steps": "1. 点击「下载视频」按钮\n2. 观察下载行为",
        "expected": "1. 按钮可正常点击\n2. 触发视频文件下载\n3. 下载文件格式正确",
        "requirement": "需求1"
    },

    # ===== 模块7: 对话交互 =====
    {
        "id": "TC-028",
        "module": "对话交互",
        "name": "对话输入框-发送消息",
        "priority": "P1",
        "precondition": "已进入任务方案确认页",
        "steps": "1. 在对话框输入文本\n2. 点击「发送」按钮\n3. 按Enter键发送\n4. 按Shift+Enter换行",
        "expected": "1. 文本正确显示在用户消息区域\n2. 点击发送正常发送消息\n3. Enter键发送，Shift+Enter换行\n4. 发送后输入框清空\n5. 消息滚动到最新位置",
        "requirement": "HTML原型-对话"
    },
    {
        "id": "TC-029",
        "module": "对话交互",
        "name": "对话框-图片附件上传",
        "priority": "P1",
        "precondition": "已进入任务方案确认页",
        "steps": "1. 点击输入框左侧附件按钮\n2. 选择一张图片上传\n3. 观察对话框中图片展示",
        "expected": "1. 文件选择器正常弹出\n2. 选择图片后提示已选择\n3. 图片可附加到消息中发送",
        "requirement": "HTML原型-对话"
    },
    {
        "id": "TC-030",
        "module": "对话交互",
        "name": "AI方案结果-对话框展示",
        "priority": "P0",
        "precondition": "已创建AI自动生成脚本任务",
        "steps": "1. 提交任务后进入方案确认页\n2. 观察对话框AI消息区域\n3. 查看视频整体方案卡片\n4. 点击「继续生成」按钮",
        "expected": "1. AI消息区域展示整体方案摘要\n2. 方案卡片包含：产品名、总时长、视频风格、钩子类型、分段规划\n3. 分段规划展示各片段主题和时长\n4. 点击「继续生成」后触发后续流程",
        "requirement": "需求1-步骤5-9"
    },

    # ===== 模块8: 后台管理-Agent提示词 =====
    {
        "id": "TC-031",
        "module": "后台管理-Agent提示词",
        "name": "后台管理-视频解析任务TAB与AI自动生成脚本TAB切换",
        "priority": "P1",
        "precondition": "1. 具有后台管理权限\n2. 进入后台管理-AI视频-AI视频提示词页面",
        "steps": "1. 查看页面TAB布局\n2. 点击「视频解析任务」TAB\n3. 点击「AI自动生成脚本任务」TAB\n4. 观察各TAB下的Agent列表",
        "expected": "1. 两个TAB正确显示\n2. 视频解析任务TAB下展示原有Agent\n3. AI自动生成脚本任务TAB下展示新增Agent\n4. TAB切换流畅，内容正确切换",
        "requirement": "需求2"
    },
    {
        "id": "TC-032",
        "module": "后台管理-Agent提示词",
        "name": "卖点提炼提示词-配置与查看",
        "priority": "P1",
        "precondition": "进入AI自动生成脚本任务TAB",
        "steps": "1. 在Agent列表中找到「卖点提炼提示词」\n2. 查看提示词配置内容\n3. 确认角色定义、需求描述、限制规则、输出格式是否完整\n4. 检查JSON输出格式定义",
        "expected": "1. 卖点提炼提示词完整展示\n2. 包含角色定义（电商运营）\n3. 包含需求描述（提炼痛点、卖点、目标人群等）\n4. 包含限制规则（违禁词过滤、字数限制）\n5. 包含JSON输出模板\n6. 提示词可编辑保存",
        "requirement": "需求2-Agent"
    },

    # ===== 边界值测试 =====
    {
        "id": "TC-033",
        "module": "边界值测试",
        "name": "视频长度-最小值输入",
        "priority": "P2",
        "precondition": "打开AI自动生成脚本任务弹窗",
        "steps": "1. 在视频长度输入框输入最小值（如：1秒）\n2. 填写其他必填项\n3. 点击提交",
        "expected": "1. 系统接受1秒的视频长度输入\n2. 任务正常创建\n3. 后续生成按1秒时长处理",
        "requirement": "需求1-步骤1a"
    },
    {
        "id": "TC-034",
        "module": "边界值测试",
        "name": "视频长度-最大值60秒",
        "priority": "P2",
        "precondition": "打开创建任务弹窗",
        "steps": "1. 在视频长度输入框输入60（最大值）\n2. 填写其他必填项\n3. 点击提交\n4. 尝试输入61",
        "expected": "1. 60秒可正常提交\n2. 视频长度提示「最长60秒」\n3. 若超过60秒应给出限制提示",
        "requirement": "视频解析任务表单提示"
    },
    {
        "id": "TC-035",
        "module": "边界值测试",
        "name": "商品名称-超长输入处理",
        "priority": "P2",
        "precondition": "打开AI自动生成脚本任务弹窗",
        "steps": "1. 在商品名称输入框输入超长文本（>100字符）\n2. 提交任务\n3. 观察任务卡片展示",
        "expected": "1. 系统接受超长名称或给出长度限制提示\n2. 任务卡片中名称正确截断显示（省略号）\n3. 名称在卖点提炼中正确传递，product_name字段≤10字",
        "requirement": "需求1-步骤1a"
    },

    # ===== 异常测试 =====
    {
        "id": "TC-036",
        "module": "异常测试",
        "name": "无效商品URL-解析失败处理",
        "priority": "P2",
        "precondition": "1. 已创建AI任务\n2. 输入了一个无效或不存在的商品URL",
        "steps": "1. 输入无效URL提交任务\n2. 等待系统解析\n3. 观察错误提示",
        "expected": "1. 系统给出友好的错误提示\n2. 错误信息明确说明原因（如：链接无效、商品不存在）\n3. 不出现系统崩溃或空白页面\n4. 用户可返回修改URL重新提交",
        "requirement": "需求1-步骤2"
    },
    {
        "id": "TC-037",
        "module": "异常测试",
        "name": "网络异常-生成过程中断处理",
        "priority": "P2",
        "precondition": "已触发视频素材图片生成",
        "steps": "1. 生成过程中模拟网络断开\n2. 观察页面反应\n3. 网络恢复后重新生成",
        "expected": "1. 生成中断时有超时或失败提示\n2. 不出现页面无响应或崩溃\n3. 页面状态可恢复\n4. 用户可重新点击生成按钮重试",
        "requirement": "需求1-步骤8-9"
    },
    {
        "id": "TC-038",
        "module": "异常测试",
        "name": "空商品名称-卖点提炼异常处理",
        "priority": "P1",
        "precondition": "已创建AI任务但商品名称为空",
        "steps": "1. 通过非常规方式提交空商品名称任务\n2. 观察卖点提炼步骤的反应\n3. 检查错误处理",
        "expected": "1. 前端校验阻止空名称提交\n2. 若绕过前端校验，后端应有兜底校验\n3. 返回明确错误提示而非系统异常",
        "requirement": "需求1-步骤3"
    },
]

wb = Workbook()

# ===== Sheet 1: 冒烟测试用例 =====
ws = wb.active
ws.title = "冒烟测试用例"

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, size=11, color="FFFFFF", name="微软雅黑")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

priority_colors = {
    "P0": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "P1": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "P2": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
}

priority_fonts = {
    "P0": Font(bold=True, size=10, color="9C0006", name="微软雅黑"),
    "P1": Font(bold=True, size=10, color="9C6500", name="微软雅黑"),
    "P2": Font(bold=True, size=10, color="006100", name="微软雅黑"),
}

module_colors = {
    "任务创建与管理": "E8F0FE",
    "商品信息处理": "F3E8FF",
    "方案生成": "FFF3E0",
    "视频素材图片生成": "E8F5E9",
    "方案内容编辑": "FCE4EC",
    "视频生成与导出": "E0F2F1",
    "对话交互": "F9FBE7",
    "后台管理-Agent提示词": "EDE7F6",
    "边界值测试": "ECEFF1",
    "异常测试": "FFEBEE",
}

headers = ["用例ID", "模块", "用例名称", "优先级", "前置条件", "测试步骤", "预期结果", "关联需求"]

for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment
    cell.border = thin_border

for row_idx, case in enumerate(test_cases, 2):
    data = [
        case["id"],
        case["module"],
        case["name"],
        case["priority"],
        case["precondition"],
        case["steps"],
        case["expected"],
        case["requirement"],
    ]
    module_fill = PatternFill(
        start_color=module_colors.get(case["module"], "FFFFFF"),
        end_color=module_colors.get(case["module"], "FFFFFF"),
        fill_type="solid"
    )
    for col_idx, value in enumerate(data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=10, name="微软雅黑")

        if col_idx == 1:
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif col_idx == 2:
            cell.fill = module_fill
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif col_idx == 4:
            cell.fill = priority_colors.get(value, None)
            cell.font = priority_fonts.get(value, Font(size=10, name="微软雅黑"))
            cell.alignment = Alignment(horizontal="center", vertical="top")

column_widths = {
    1: 10, 2: 16, 3: 30, 4: 8, 5: 30, 6: 45, 7: 40, 8: 18
}
for col_idx, width in column_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 30
for row_idx in range(2, len(test_cases) + 2):
    ws.row_dimensions[row_idx].height = 80

ws.auto_filter.ref = f"A1:H{len(test_cases) + 1}"
ws.freeze_panes = "A2"

# ===== Sheet 2: 测试用例统计 =====
ws2 = wb.create_sheet(title="测试用例统计")

stat_header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
stat_header_font = Font(bold=True, size=12, color="FFFFFF", name="微软雅黑")
section_font = Font(bold=True, size=11, color="4472C4", name="微软雅黑")
section_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

ws2.column_dimensions['A'].width = 28
ws2.column_dimensions['B'].width = 12
ws2.column_dimensions['C'].width = 12
ws2.column_dimensions['D'].width = 12
ws2.column_dimensions['E'].width = 12

ws2.merge_cells('A1:E1')
title_cell = ws2.cell(row=1, column=1, value="AI视频工具3.1 - 分镜图生成视频 - 冒烟测试用例统计")
title_cell.font = Font(bold=True, size=14, color="1A237E", name="微软雅黑")
title_cell.alignment = Alignment(horizontal="center", vertical="center")
title_cell.fill = PatternFill(start_color="E8EAF6", end_color="E8EAF6", fill_type="solid")
ws2.row_dimensions[1].height = 36

row = 3
ws2.cell(row=row, column=1, value="一、模块分布统计").font = section_font
ws2.cell(row=row, column=1).fill = section_fill
for c in range(1, 6):
    ws2.cell(row=row, column=c).fill = section_fill
    ws2.cell(row=row, column=c).border = thin_border

row += 1
stat_headers = ["模块名称", "P0数量", "P1数量", "P2数量", "小计"]
for col_idx, h in enumerate(stat_headers, 1):
    cell = ws2.cell(row=row, column=col_idx, value=h)
    cell.fill = stat_header_fill
    cell.font = stat_header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border
ws2.row_dimensions[row].height = 26

module_order = ["任务创建与管理", "商品信息处理", "方案生成", "视频素材图片生成",
                "方案内容编辑", "视频生成与导出", "对话交互", "后台管理-Agent提示词",
                "边界值测试", "异常测试"]
module_data = {}
for tc in test_cases:
    m = tc["module"]
    if m not in module_data:
        module_data[m] = {"P0": 0, "P1": 0, "P2": 0}
    module_data[m][tc["priority"]] += 1

total_p0 = total_p1 = total_p2 = 0
row += 1
for mod in module_order:
    if mod in module_data:
        d = module_data[mod]
        vals = [mod, d["P0"], d["P1"], d["P2"], d["P0"] + d["P1"] + d["P2"]]
        total_p0 += d["P0"]
        total_p1 += d["P1"]
        total_p2 += d["P2"]
        for col_idx, v in enumerate(vals, 1):
            cell = ws2.cell(row=row, column=col_idx, value=v)
            cell.border = thin_border
            cell.font = Font(size=10, name="微软雅黑")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[row].height = 24
        row += 1

row_data = ["合计", total_p0, total_p1, total_p2, total_p0 + total_p1 + total_p2]
for col_idx, v in enumerate(row_data, 1):
    cell = ws2.cell(row=row, column=col_idx, value=v)
    cell.border = thin_border
    cell.font = Font(bold=True, size=11, name="微软雅黑")
    cell.fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[row].height = 26

row += 2
ws2.cell(row=row, column=1, value="二、优先级分布统计").font = section_font
for c in range(1, 6):
    ws2.cell(row=row, column=c).fill = section_fill
    ws2.cell(row=row, column=c).border = thin_border

row += 1
pri_headers = ["优先级", "用例数量", "占比", "说明", ""]
for col_idx, h in enumerate(pri_headers, 1):
    cell = ws2.cell(row=row, column=col_idx, value=h)
    cell.fill = stat_header_fill
    cell.font = stat_header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border
ws2.row_dimensions[row].height = 26

pri_data = [
    ("P0 - 核心流程", total_p0, f"{total_p0/len(test_cases)*100:.1f}%", "核心业务流程，必须通过", "FFC7CE"),
    ("P1 - 重要功能", total_p1, f"{total_p1/len(test_cases)*100:.1f}%", "重要功能/校验，应通过", "FFEB9C"),
    ("P2 - 边界与异常", total_p2, f"{total_p2/len(test_cases)*100:.1f}%", "边界条件/异常处理，建议通过", "C6EFCE"),
]

row += 1
for p_name, p_count, p_ratio, p_desc, p_color in pri_data:
    vals = [p_name, p_count, p_ratio, p_desc]
    for col_idx, v in enumerate(vals, 1):
        cell = ws2.cell(row=row, column=col_idx, value=v)
        cell.border = thin_border
        cell.font = Font(size=10, name="微软雅黑")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill(start_color=p_color, end_color=p_color, fill_type="solid")
    ws2.row_dimensions[row].height = 28
    row += 1

row += 1
ws2.merge_cells(f'A{row}:E{row}')
cell = ws2.cell(row=row, column=1, value=f"总计：{len(test_cases)} 条冒烟测试用例")
cell.font = Font(bold=True, size=12, color="1A237E", name="微软雅黑")
cell.alignment = Alignment(horizontal="center", vertical="center")
cell.border = thin_border
ws2.row_dimensions[row].height = 30

row += 2
ws2.cell(row=row, column=1, value="三、测试覆盖说明").font = section_font
for c in range(1, 6):
    ws2.cell(row=row, column=c).fill = section_fill
    ws2.cell(row=row, column=c).border = thin_border

row += 1
info_items = [
    "测试范围：AI视频工具3.1 - 分镜图生成视频功能（一期）",
    "需求版本：AI视频工具3.1",
    f"设计日期：2026-06-01",
    "覆盖功能模块：10个",
    "主要流程覆盖：全覆盖（任务创建 → 商品解析 → 卖点提炼 → 方案生成 → 素材生成 → 视频导出）",
    "一期未覆盖项：视频素材图片的对话交互功能（明确标注一期不实现）",
    "未明确需求说明：Agent提示词中除「卖点提炼提示词」外，其余Agent提示词内容待补充",
]
for item in info_items:
    cell = ws2.cell(row=row, column=1, value=item)
    ws2.merge_cells(f'A{row}:E{row}')
    cell.font = Font(size=10, name="微软雅黑")
    cell.alignment = Alignment(vertical="center")
    cell.border = thin_border
    ws2.row_dimensions[row].height = 22
    row += 1

output_path = "/Users/lanwang/Documents/trae_projects/new World/AI视频工具3.1_分镜图生成视频_冒烟测试用例.xlsx"
wb.save(output_path)
print(f"✅ 冒烟测试用例Excel已生成：{output_path}")
print(f"   共 {len(test_cases)} 条测试用例")
print(f"   P0: {total_p0} 条 | P1: {total_p1} 条 | P2: {total_p2} 条")
